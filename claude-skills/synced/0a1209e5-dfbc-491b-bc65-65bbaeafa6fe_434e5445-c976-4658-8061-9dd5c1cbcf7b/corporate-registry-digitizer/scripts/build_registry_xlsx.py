#!/usr/bin/env python3
"""
Builds the formatted strain registry xlsx file from a JSON list of rows.

Usage:
    python build_registry_xlsx.py --input rows.json --output Strains_updated.xlsx [--title "Strain Registry"]

rows.json is a list of objects. Each object may contain any of these keys
(missing keys are left blank in the sheet):

    code, parental_strain, plasmid, protein, resistance, clon,
    lnb, project, snapgene_file, date, notes

"notes" is where flags like "POSSIBLE DUPLICATE OF ROW X",
"POSSIBLE STRIKETHROUGH, PLEASE VERIFY", "CONTINUES ON NEXT PAGE",
or spelling correction notes should go.
"""

import argparse
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VIBRANT_ORANGE = "FF6600"
ELECTRIC_BLUE = "005CE6"

COLUMNS = [
    ("code", "Code"),
    ("parental_strain", "Parental strain"),
    ("plasmid", "Plasmid"),
    ("protein", "Protein"),
    ("resistance", "Resistance"),
    ("clon", "Clon"),
    ("lnb", "LNB"),
    ("project", "Project"),
    ("snapgene_file", "SnapGene file?"),
    ("date", "Date"),
    ("notes", "Notes"),
]


def build_workbook(rows, title=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Registry"

    start_row = 1

    if title:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
        title_cell = ws.cell(row=1, column=1, value=title)
        title_cell.font = Font(bold=True, size=14, color="FFFFFF")
        title_cell.fill = PatternFill(start_color=VIBRANT_ORANGE, end_color=VIBRANT_ORANGE, fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24
        start_row = 2

    header_row = start_row
    for col_index, (_, label) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_index, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=ELECTRIC_BLUE, end_color=ELECTRIC_BLUE, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_offset, row_data in enumerate(rows, start=1):
        excel_row = header_row + row_offset
        for col_index, (key, _) in enumerate(COLUMNS, start=1):
            value = row_data.get(key, "")
            ws.cell(row=excel_row, column=col_index, value=value)

    widths = [16, 20, 20, 20, 16, 10, 14, 20, 16, 12, 34]
    for col_index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_index)].width = width

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate

    return wb


def main():
    parser = argparse.ArgumentParser(description="Build the formatted strain registry xlsx file.")
    parser.add_argument("--input", required=True, help="Path to a JSON file containing a list of row objects.")
    parser.add_argument("--output", required=True, help="Path to write the resulting xlsx file to.")
    parser.add_argument("--title", default=None, help="Optional title placed in a merged, colored row above the headers.")
    args = parser.parse_args()

    with open(args.input, "r", encoding="utf-8") as f:
        rows = json.load(f)

    if not isinstance(rows, list):
        print("Input JSON must be a list of row objects.", file=sys.stderr)
        sys.exit(1)

    wb = build_workbook(rows, title=args.title)
    wb.save(args.output)
    print(f"Wrote {len(rows)} rows to {args.output}")


if __name__ == "__main__":
    main()
